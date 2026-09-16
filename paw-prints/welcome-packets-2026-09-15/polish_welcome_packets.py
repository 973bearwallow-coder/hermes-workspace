from __future__ import annotations

import shutil
import zipfile
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
import lxml.etree as etree

ROOT = Path(__file__).resolve().parent
ORIGINALS = ROOT / "originals"
FINAL = ROOT / "final"

DEEP_GREEN = "315C3B"
BRAND_GREEN = "4F7D3A"
PALE_GREEN = "EAF2E3"
PALE_GOLD = "FFF1C7"
WARM_CREAM = "FFF9E8"
DARK_BROWN = "3D2E25"
MUTED = "6B675F"
WHITE = "FFFFFF"
BORDER = "B8C4B1"

TITLE_PHRASES = (
    "Service Agreement",
    "Veterinary Release Agreement",
    "Pet Lodging Addendum",
    "Key Handling Agreement",
    "Contact Information",
    "Pet Information Form",
    "Service Request",
)

MAJOR_HEADINGS = {
    "New Client Packet",
    "Suggested Leave-Out Checklist",
    "Paw Prints Pet Services LLC",
    "Service Agreement",
    "Veterinary Release Agreement",
    "Pet Lodging Addendum",
    "Key Handling Agreement",
    "Contact Information",
    "Pet Information Form",
    "Service Request",
    "Important Terms",
}

MINOR_HEADINGS = {
    "The New Client Packet Includes:",
    "At the initial interview (meet and greet):",
    "If applicable:",
    "Keys:",
    "Contact Us:",
    "Key Returns",
    "Emergency contacts",
    "Things to bring for pet lodging or daycare:",
    "Schedule",
    "Tasks",
    "Additional Comments:",
}


def iter_paragraphs(parent) -> Iterable:
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        element.set(qn("w:space"), "0")


def color_runs(paragraph, color: str, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            run.bold = bold


def normalized(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def title_table(table) -> bool:
    text = normalized(" ".join(cell.text for row in table.rows for cell in row.cells))
    return any(phrase in text for phrase in TITLE_PHRASES)


def rates_table(table) -> bool:
    text = normalized(" ".join(cell.text for row in table.rows for cell in row.cells)).upper()
    return "VISIT TYPE" in text and "RATE" in text


def map_existing_fill(fill: str, cell_text: str) -> str:
    fill = fill.upper().replace("#", "")
    text = normalized(cell_text).lower()
    if "owner" in text:
        return PALE_GREEN
    if text == "pet" or text.startswith("pet "):
        return PALE_GOLD
    if fill in {"00FF00", "92D050", "70AD47", "C6E0B4", "A9D18E"}:
        return PALE_GREEN
    if fill in {"FFFF00", "FFF2CC", "FFE699", "FFD966"}:
        return PALE_GOLD
    if fill in {"F4B183", "F8CBAD", "FCE4D6", "E2F0D9", "DDEBF7", "D9EAD3", "E4DFEC", "EADCF8"}:
        return WARM_CREAM
    return fill


def polish_table(table) -> None:
    if rates_table(table):
        return
    is_title = title_table(table)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = tc_pr.find(qn("w:shd"))
            if shd is not None:
                current = shd.get(qn("w:fill"), "FFFFFF")
                mapped = map_existing_fill(current, cell.text)
                shd.set(qn("w:fill"), mapped)
                shd.set(qn("w:val"), "clear")
            set_cell_borders(cell)
            if is_title and c_idx == 0:
                set_cell_fill(cell, PALE_GREEN)
                for paragraph in cell.paragraphs:
                    color_runs(paragraph, DEEP_GREEN, True)
            for paragraph in cell.paragraphs:
                text = normalized(paragraph.text)
                if text in MAJOR_HEADINGS:
                    color_runs(paragraph, DEEP_GREEN, True)
                elif text in MINOR_HEADINGS:
                    color_runs(paragraph, BRAND_GREEN, True)


def polish_paragraph(paragraph) -> None:
    text = normalized(paragraph.text)
    if not text:
        return
    if text in MAJOR_HEADINGS:
        color_runs(paragraph, DEEP_GREEN, True)
        paragraph.paragraph_format.keep_with_next = True
    elif text in MINOR_HEADINGS or (len(text) <= 55 and text.endswith(":")):
        color_runs(paragraph, BRAND_GREEN, True)
        paragraph.paragraph_format.keep_with_next = True
    elif text.startswith(("A. ", "B. ", "C. ", "D. ", "E. ", "F. ", "G. ", "H. ")):
        color_runs(paragraph, DEEP_GREEN, None)
    elif "Thank you for choosing Paw Prints Pet Services" in text or "We look forward to providing" in text:
        color_runs(paragraph, BRAND_GREEN, True)


def polish_header_footer(section) -> None:
    for part in (section.header, section.footer):
        for paragraph in iter_paragraphs(part):
            if paragraph.text.strip():
                color_runs(paragraph, MUTED, None)


def document_text_parts(path: Path) -> dict[str, list[str]]:
    parts = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            data = archive.read(name)
            try:
                root = etree.fromstring(data)
                texts = root.xpath("//w:t/text()", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
            except Exception:
                continue
            if texts:
                parts[name] = texts
    return parts


def polish(source: Path, output: Path) -> None:
    before = document_text_parts(source)
    document = Document(str(source))
    document.core_properties.title = source.stem.replace("-", " — ")
    document.core_properties.subject = "Paw Prints Pet Services new-client welcome packet"
    document.core_properties.keywords = "Paw Prints Pet Services, welcome packet, client forms"
    for paragraph in document.paragraphs:
        polish_paragraph(paragraph)
    for table in document.tables:
        polish_table(table)
    for section in document.sections:
        polish_header_footer(section)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    after = document_text_parts(output)
    if before != after:
        changed = sorted(set(before) | set(after))
        mismatches = [name for name in changed if before.get(name) != after.get(name)]
        raise RuntimeError(f"Text changed while polishing {source.name}: {mismatches}")


def main() -> None:
    sources = [
        ORIGINALS / "Paw Prints-Welcome packet.docx",
        ORIGINALS / "Paw Prints-Welcome packet with Pet Lodging.docx",
    ]
    for source in sources:
        output = FINAL / source.name.replace(".docx", " - Professional Refresh.docx")
        polish(source, output)
        print(output)


if __name__ == "__main__":
    main()
