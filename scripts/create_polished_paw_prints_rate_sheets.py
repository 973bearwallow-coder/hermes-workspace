#!/usr/bin/env python3
"""Create polished Paw Prints rates-and-services Word documents."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('/home/tom/hermes-workspace/inbox/janie-rate-sheets')
OUT = ROOT / 'polished'
OUT.mkdir(parents=True, exist_ok=True)
LOGO = ROOT / 'inspection/Paw Prints Pet Services Rates 2026-images/image1.png'

GREEN = '507A45'
DARK_GREEN = '31543A'
TAN = 'EDE3D3'
LIGHT_GREEN = 'EAF1E7'
CREAM = 'FBF8F2'
BROWN = '8A5A2B'
WHITE = 'FFFFFF'
CHARCOAL = '2E342F'
GRID = 'B7A98F'

RATES_COMMON = [
    ('Initial Consultation', 'FREE'),
    ('Regular Visit', '$27'),
    ('Express Visit', '$21'),
    ('45-Minute Visit', '$40'),
    ('1-Hour Visit', '$50'),
    ('Daycare', '$45 per dog'),
    ('Drop-In', '$10'),
    ('Transportation', '$20'),
]
FEES = [
    ('Returned check', '$20'),
    ('Payment 30+ days late', '$15'),
    ('Federal holidays', '+$5 per visit'),
    ('Before 8 AM / after 8 PM', '+$5 per visit'),
]
SERVICES = [
    ('Initial Booking Consultation', '30–60 minutes',
     'Complete paperwork, transfer keys, answer questions, meet your pets, and review detailed care instructions.'),
    ('Regular Visit', '25–30 minutes',
     'Our most popular option—time for a good walk and/or play, plus plenty of personal attention.'),
    ('Express Visit', '10–15 minutes',
     'Ideal for a quick potty break, short walk, litter-box care, an easy-keeper visit, or a third feeding.'),
    ('Drop-In', '3 minutes or less',
     'A brief stop to exchange keys or handle a quick household check, such as an iron, stove, sprinkler, or window.'),
    ('Pet Lodging (Boarding)', 'By arrangement',
     'Your dog stays in your sitter’s home with feeding, walks, attentive care, and extra TLC. Available case by case.'),
    ('Daycare', '8 AM–8 PM',
     'Daytime care in your sitter’s home. Available case by case.'),
    ('Transportation', 'Within 5 miles',
     'Pickup or drop-off for lodging, daycare, or another nearby destination.'),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def borders(cell, color=GRID, size='5'):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_keep(paragraph, keep_next=False):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement('w:keepLines')
    p_pr.append(keep_lines)
    if keep_next:
        node = OxmlElement('w:keepNext')
        p_pr.append(node)


def add_text(p, text, size=8.4, bold=False, color=CHARCOAL, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def configure_document(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.3)
    sec.bottom_margin = Inches(0.3)
    sec.left_margin = Inches(0.42)
    sec.right_margin = Inches(0.42)

    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(8.4)
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def add_header(doc, legacy, effective):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(5.7)
    table.columns[2].width = Inches(2.1)
    for c in table.rows[0].cells:
        set_cell_margins(c, 10, 50, 10, 50)
        borders(c, WHITE, '0')

    if LOGO.exists():
        p = table.cell(0, 0).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Inches(1.72))

    p = table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    add_text(p, 'RATES & SERVICES', 20, True, DARK_GREEN)
    if legacy:
        p2 = table.cell(0, 1).add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p2, 'LEGACY CLIENT PRICING', 9.5, True, BROWN)

    p = table.cell(0, 2).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, '703-244-7390', 10.5, True, DARK_GREEN)
    p2 = table.cell(0, 2).add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p2, 'pawprintsva.com', 9.5, True, BROWN)
    p3 = table.cell(0, 2).add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p3, f'Effective {effective}', 8, False, CHARCOAL)

    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.cell(0, 0)
    shade(c, GREEN)
    borders(c, GREEN)
    set_cell_margins(c, 55, 90, 55, 90)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, 'Trusted, personalized pet care in Falls Church, Virginia', 9.3, True, WHITE)


def add_section_title(cell, title):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    set_keep(p, True)
    add_text(p, title.upper(), 10.5, True, DARK_GREEN)
    return p


def add_rate_table(cell, lodging_rate):
    add_section_title(cell, 'Visit Rates')
    rows = [('Pet Lodging*', lodging_rate + ' per dog, per night')] + RATES_COMMON[5:]
    rows = RATES_COMMON[:5] + rows
    table = cell.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.88)
    table.columns[1].width = Inches(1.48)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for idx, label in enumerate(('SERVICE', 'RATE')):
        shade(hdr[idx], GREEN)
        borders(hdr[idx], GREEN)
        set_cell_margins(hdr[idx], 55, 70, 55, 70)
        p = hdr[idx].paragraphs[0]
        add_text(p, label, 8, True, WHITE)
    for i, (name, rate) in enumerate(rows):
        cells = table.add_row().cells
        fill = CREAM if i % 2 == 0 else WHITE
        for c in cells:
            shade(c, fill)
            borders(c)
            set_cell_margins(c, 45, 70, 45, 70)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(cells[0].paragraphs[0], name, 8.2, True if name == 'Pet Lodging*' else False)
        add_text(cells[1].paragraphs[0], rate, 8.2, True if name == 'Pet Lodging*' else False, BROWN if name == 'Pet Lodging*' else CHARCOAL)

    add_section_title(cell, 'Additional Fees')
    fee_table = cell.add_table(rows=0, cols=2)
    fee_table.autofit = False
    fee_table.columns[0].width = Inches(1.88)
    fee_table.columns[1].width = Inches(1.48)
    for i, (name, rate) in enumerate(FEES):
        cells = fee_table.add_row().cells
        fill = LIGHT_GREEN if i % 2 == 0 else WHITE
        for c in cells:
            shade(c, fill)
            borders(c)
            set_cell_margins(c, 43, 70, 43, 70)
        add_text(cells[0].paragraphs[0], name, 7.8)
        add_text(cells[1].paragraphs[0], rate, 7.8, True)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    add_text(p, 'GOOD TO KNOW', 8.2, True, BROWN)
    notes = [
        'No additional charge for multiple pets during dog walking, pet sitting, or cat-care visits.',
        '* Lodging is priced per pet, per night. Add $10 per night for puppies or senior dogs needing extra care or cleanup.',
        'Visits are scheduled within a preferred three-hour service window.',
    ]
    for note in notes:
        p = cell.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.13)
        p.paragraph_format.first_line_indent = Inches(-0.08)
        p.paragraph_format.space_after = Pt(1)
        add_text(p, note, 7.35)


def add_services(cell):
    add_section_title(cell, 'Services at a Glance')
    for idx, (name, time, desc) in enumerate(SERVICES):
        block = cell.add_table(rows=1, cols=1)
        block.autofit = False
        shade(block.cell(0, 0), LIGHT_GREEN if idx % 2 == 0 else CREAM)
        borders(block.cell(0, 0), WHITE, '0')
        set_cell_margins(block.cell(0, 0), 35, 80, 35, 80)
        p = block.cell(0, 0).paragraphs[0]
        set_keep(p)
        add_text(p, name, 8.5, True, DARK_GREEN)
        add_text(p, f'  •  {time}', 7.7, True, BROWN)
        p2 = block.cell(0, 0).add_paragraph()
        add_text(p2, desc, 7.45)


def add_terms(cell):
    add_section_title(cell, 'Policies & Payment')
    policies = [
        ('Scheduling', 'Please reserve enough time for the care requested. If additional time is needed, it will be added and billed accordingly.'),
        ('Payment', 'Payment is due upon receipt of the invoice. Cash is preferred; checks, PayPal, and Zelle are also accepted.'),
        ('Cancellations', 'Please cancel by 9 AM on the day of service. Later cancellations are charged in full.'),
        ('Holidays', 'Federal holidays include an additional fee and may have limited availability.'),
    ]
    for title, text in policies:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        set_keep(p)
        add_text(p, title + ': ', 8.1, True, BROWN)
        add_text(p, text, 7.7)

    callout = cell.add_table(rows=1, cols=1)
    c = callout.cell(0, 0)
    shade(c, GREEN)
    borders(c, GREEN)
    set_cell_margins(c, 75, 90, 75, 90)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, 'Happy pets.\nPeace of mind.', 8.2, True, WHITE, True)


def add_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, 'Paw Prints Pet Services LLC  •  703-244-7390  •  pawprintsva.com  •  Rates subject to change', 7.3, False, DARK_GREEN)


def build(filename, lodging_rate, legacy):
    doc = Document()
    configure_document(doc)
    add_header(doc, legacy, 'November 2026')

    outer = doc.add_table(rows=1, cols=3)
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer.autofit = False
    outer.columns[0].width = Inches(3.45)
    outer.columns[1].width = Inches(4.0)
    outer.columns[2].width = Inches(2.25)
    for c in outer.rows[0].cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(c, 30, 90, 20, 90)
        borders(c, WHITE, '0')

    add_rate_table(outer.cell(0, 0), lodging_rate)
    add_services(outer.cell(0, 1))
    add_terms(outer.cell(0, 2))
    add_footer(doc)

    path = OUT / filename
    doc.save(path)
    return path


if __name__ == '__main__':
    paths = [
        build('Paw Prints Rates & Services 2026 - Legacy 1 - Polished.docx', '$80', True),
        build('Paw Prints Rates & Services 2026 - Legacy 2 - Polished.docx', '$80', True),
        build('Paw Prints Rates & Services 2026 - Polished.docx', '$100', False),
    ]
    for path in paths:
        print(path)
