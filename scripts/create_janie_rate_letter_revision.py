#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

out = Path('/home/tom/hermes-workspace/inbox/janie-rate-adjustment/Paw Prints price change, November 2026 - Atlas suggested revision.docx')
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08

p = doc.add_paragraph('September 1, 2026')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph('Dear Wonderful Pet Care Clients,')

doc.add_paragraph(
    'Thank you for trusting me as part of your pet care team. I am grateful for the opportunity '
    'to care for your pets and for the confidence you place in Paw Prints Pet Services.'
)

doc.add_paragraph(
    'To continue providing the reliable, professional, and loving care you have come to count on, '
    'some of my rates will change beginning November 1, 2026.'
)

p = doc.add_paragraph()
r = p.add_run('New rates effective November 1, 2026')
r.bold = True
for text in [
    '$27 for a 30-minute visit',
    '$21 for a 15-minute visit',
    '$80 per night, per dog, for pet lodging',
]:
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.25)

p = doc.add_paragraph()
r = p.add_run('Longtime-client discount: ')
r.bold = True
p.add_run(
    'Your legacy discount will remain in place: $2 off each walk and '
    '$20 off each night of pet lodging.'
)

doc.add_paragraph(
    'Please see the attached pricing sheet for additional services and pricing options.'
)

doc.add_paragraph(
    'Paw Prints Pet Services is a family-owned small business, fully insured and backed by more '
    'than 18 years of experience. I remain committed to keeping your pets loved, safe, and happy. '
    'Our home-style dog lodging offers a cozy family atmosphere, individualized attention, and a '
    'smaller, controlled setting—all at competitive rates compared with commercial kennels and '
    'daycare facilities. It is truly a home away from home.'
)

doc.add_paragraph(
    'I sincerely appreciate your loyalty and understanding. Please contact me if you have any questions.'
)

doc.add_paragraph('Warmly,')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run('Jane Torok')
r.bold = True
doc.add_paragraph('Paw Prints Pet Services, LLC')
doc.add_paragraph('703-244-7390')
doc.add_paragraph('www.pawprintsva.com')

doc.save(out)
print(out)
